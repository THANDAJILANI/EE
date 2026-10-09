import io
import csv
import json
from typing import Dict, Any, List
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from app.models.paper import ResearchPaper
from app.models.literature import LiteratureReview
from app.models.comparison import PaperComparison

def generate_paper_pdf(paper: ResearchPaper, export_type: str = "executive_summary") -> io.BytesIO:
    """
    Generates a publication-grade PDF report using ReportLab.
    Supports 'executive_summary' and 'full_analysis'.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=45
    )
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#1e3a8a"),
        spaceAfter=10
    )
    h2_style = ParagraphStyle(
        'H2',
        parent=styles['Heading2'],
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#0f172a"),
        spaceBefore=12,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#334155")
    )
    meta_style = ParagraphStyle(
        'Meta',
        parent=styles['Italic'],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#64748b")
    )

    story = []

    # Title & Branding
    story.append(Paragraph("ResearchMate AI — Academic Paper Report", meta_style))
    story.append(Spacer(1, 5))
    story.append(Paragraph(paper.title, title_style))

    meta = paper.metadata_rel
    authors_str = ", ".join(meta.authors) if (meta and meta.authors) else "Authors not reported"
    year_str = str(meta.publication_year) if (meta and meta.publication_year) else "Year n.d."
    journal_str = meta.journal_or_conference if (meta and meta.journal_or_conference) else "N/A"
    story.append(Paragraph(f"<b>Authors:</b> {authors_str} &nbsp;|&nbsp; <b>Year:</b> {year_str} &nbsp;|&nbsp; <b>Venue:</b> {journal_str}", meta_style))
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceAfter=12))

    analysis = paper.analysis
    if not analysis:
        story.append(Paragraph("Analysis not yet generated for this document.", body_style))
        doc.build(story)
        buffer.seek(0)
        return buffer

    # Executive Summary Section
    story.append(Paragraph("1. Executive Summary", h2_style))
    summary_text = analysis.detailed_summary or analysis.concise_summary or "Summary not available."
    story.append(Paragraph(summary_text, body_style))
    story.append(Spacer(1, 8))

    if analysis.main_contribution:
        story.append(Paragraph(f"<b>Main Contribution:</b> {analysis.main_contribution}", body_style))
        story.append(Spacer(1, 4))
    if analysis.why_it_matters:
        story.append(Paragraph(f"<b>Why It Matters:</b> {analysis.why_it_matters}", body_style))
        story.append(Spacer(1, 8))

    if export_type == "full_analysis":
        # Research Objectives
        story.append(Paragraph("2. Research Objectives", h2_style))
        objs = analysis.research_objectives or {}
        story.append(Paragraph(f"<b>Primary Objective:</b> {objs.get('primary_objective', 'Not reported in the paper')}", body_style))
        story.append(Spacer(1, 4))
        story.append(Paragraph(f"<b>Problem Addressed:</b> {objs.get('problem_addressed', 'Not reported in the paper')}", body_style))
        story.append(Spacer(1, 4))
        story.append(Paragraph(f"<b>Research Gap:</b> {objs.get('research_gap', 'Not reported in the paper')}", body_style))
        story.append(Spacer(1, 8))

        # Methodology
        story.append(Paragraph("3. Methodology & Evaluation Setup", h2_style))
        meth = analysis.methodology or {}
        ds = meth.get("dataset", {})
        ds_name = ds.get("name", "Not reported in the paper") if isinstance(ds, dict) else "Not reported in the paper"
        algos = ", ".join(meth.get("algorithms_models", [])) or "Not reported in the paper"
        metrics = ", ".join(meth.get("evaluation_metrics", [])) or "Not reported in the paper"
        story.append(Paragraph(f"<b>Research Design:</b> {meth.get('research_design', 'Not reported in the paper')}", body_style))
        story.append(Paragraph(f"<b>Dataset:</b> {ds_name} &nbsp;|&nbsp; <b>Algorithms:</b> {algos}", body_style))
        story.append(Paragraph(f"<b>Evaluation Metrics:</b> {metrics}", body_style))
        story.append(Spacer(1, 8))

        # Results & Findings
        story.append(Paragraph("4. Key Results & Findings", h2_style))
        results = analysis.results_findings or {}
        for finding in results.get("main_findings", [])[:4]:
            story.append(Paragraph(f"• {finding}", body_style))
        story.append(Spacer(1, 8))

        # Limitations (Strict Separation)
        story.append(Paragraph("5. Limitations", h2_style))
        lims = analysis.limitations or {}
        story.append(Paragraph("<b>Author-Stated Limitations:</b>", body_style))
        for auth_lim in lims.get("author_stated", [])[:3]:
            story.append(Paragraph(f"• {auth_lim}", body_style))
        story.append(Spacer(1, 4))
        story.append(Paragraph("<b>AI-Inferred Potential Limitations:</b>", body_style))
        for ai_lim in lims.get("ai_inferred", [])[:2]:
            if isinstance(ai_lim, dict):
                story.append(Paragraph(f"• {ai_lim.get('limitation', '')} <i>(Evidence: {ai_lim.get('evidence_rationale', '')})</i>", body_style))
        story.append(Spacer(1, 8))

    doc.build(story)
    buffer.seek(0)
    return buffer

def generate_docx_literature_review(review: LiteratureReview) -> io.BytesIO:
    """
    Generates a structured Word document (.docx) for literature review draft.
    """
    doc = Document()
    doc.add_heading(review.title, level=0)

    p_meta = doc.add_paragraph()
    p_meta.add_run(f"Research Topic: {review.topic}\n").bold = True
    p_meta.add_run(f"Citation Style: {review.citation_style} | Generated via ResearchMate AI\n")

    # Introduction
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        f"This literature review systematically examines existing empirical and theoretical works "
        f"addressing '{review.topic}'. Synthesizing {len(review.paper_ids or [])} primary research papers, "
        f"the review establishes the state of the art, methodological trends, and recurring research gaps."
    )

    # Matrix Table
    if review.matrix_data:
        doc.add_heading("2. Literature Review Matrix", level=1)
        matrix = review.matrix_data
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Title & Year"
        hdr_cells[1].text = "Dataset"
        hdr_cells[2].text = "Methodology"
        hdr_cells[3].text = "Key Findings"
        hdr_cells[4].text = "Limitations"

        for row_data in matrix:
            row_cells = table.add_row().cells
            row_cells[0].text = f"{row_data.get('title', '')} ({row_data.get('year', '')})"
            row_cells[1].text = row_data.get('dataset', '')
            row_cells[2].text = row_data.get('methodology', '')
            row_cells[3].text = row_data.get('key_findings', '')
            row_cells[4].text = row_data.get('limitations', '')

    # Individual Paper Notes
    if review.individual_paper_notes:
        doc.add_heading("3. Individual Paper Notes", level=1)
        for note in review.individual_paper_notes:
            doc.add_heading(note.get("paper_title", "Untitled Paper"), level=2)
            p = doc.add_paragraph()
            p.add_run("Citation: ").bold = True
            p.add_run(f"{note.get('authors_year', '')}\n")
            p.add_run("Objective: ").bold = True
            p.add_run(f"{note.get('objectives', '')}\n")
            p.add_run("Method & Algorithms: ").bold = True
            p.add_run(f"{note.get('methodology', '')} - {note.get('algorithms', '')}\n")
            p.add_run("Key Results: ").bold = True
            p.add_run(f"{note.get('key_results', '')}\n")
            p.add_run("Identified Gap: ").bold = True
            p.add_run(f"{note.get('research_gap', '')}\n")
            p.add_run("Relevance to Topic: ").bold = True
            p.add_run(f"{note.get('relevance_to_topic', '')}\n")

    # Thematic Synthesis
    if review.thematic_review:
        doc.add_heading("4. Thematic Discussion", level=1)
        doc.add_paragraph(review.thematic_review)

    # Methodological Comparison
    if review.methodology_comparison:
        doc.add_heading("5. Methodological Analysis", level=1)
        doc.add_paragraph(review.methodology_comparison)

    # Gaps & Future Directions
    if review.common_gaps:
        doc.add_heading("6. Common Research Gaps", level=1)
        for g in review.common_gaps:
            doc.add_paragraph(g, style='List Bullet')

    if review.future_directions:
        doc.add_heading("7. Future Research Trajectories", level=1)
        for f in review.future_directions:
            doc.add_paragraph(f, style='List Bullet')

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def generate_csv_matrix(matrix_data: List[Dict[str, Any]]) -> str:
    """Generates a CSV string representation of the comparison or literature matrix."""
    if not matrix_data:
        return "No data available"

    output = io.StringIO()
    headers = list(matrix_data[0].keys())
    writer = csv.DictWriter(output, fieldnames=headers)
    writer.writeheader()
    for row in matrix_data:
        writer.writerow(row)
    return output.getvalue()
